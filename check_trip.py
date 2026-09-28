"""
check_trip.py
-------------
ตรวจว่าเที่ยวที่แดชบอร์ดแสดง มีอยู่จริงในไฟล์ Excel กล้วยไม้หรือไม่ และอยู่ตรงไหน

วิธีใช้ (รันในโฟลเดอร์เดียวกับ app.py):
    python check_trip.py                          # ค้นค่าตั้งต้นด้านล่าง
    python check_trip.py 6260311415403 158945     # ระบุเลขที่จะค้นเอง

สิ่งที่ทำ:
    1) เปิดไฟล์ Excel กล้วยไม้โดยตรง แล้วค้นทุกชีต ทุกช่อง (ทั้งตัวเลขและข้อความ)
    2) ดูในฐานข้อมูล dashboard.duckdb ว่าแถวนี้มาจากไฟล์/ชีต/แถวไหน และมีค่าอะไรบ้าง
"""

import re
import sys
import unicodedata
from pathlib import Path

DATA_FOLDER = Path("data")
DB_PATH = Path("dashboard.duckdb")
SOURCE_KEYWORD = "กล้วยไม้"

DEFAULT_TARGETS = ["6260311415403", "158945"]


def _key(text) -> str:
    return unicodedata.normalize("NFC", str(text)).casefold()


def _digits(value) -> str:
    """แปลงค่าเป็นตัวเลขล้วนสำหรับเทียบ: 503,046.00 -> 503046 / 6.250911179812E+12 -> 6250911179812"""
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.2f}".rstrip("0").rstrip(".").replace(".", "")
    if isinstance(value, int):
        return str(value)
    text = str(value).replace(",", "").strip()
    try:
        number = float(text)
        if number.is_integer():
            return str(int(number))
    except ValueError:
        pass
    return re.sub(r"\s+", "", text)


def find_source_file():
    files = [
        f for f in DATA_FOLDER.iterdir()
        if f.is_file() and f.suffix.lower() in {".xlsx", ".xlsm"}
        and not f.name.startswith("~$") and _key(SOURCE_KEYWORD) in _key(f.stem)
    ] if DATA_FOLDER.exists() else []
    return max(files, key=lambda f: f.stat().st_mtime_ns) if files else None


def search_excel(path: Path, targets):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    print(f"\n[1] ค้นในไฟล์ Excel: {path}")
    print("    (อ่านค่าที่คำนวณแล้วของทุกช่อง อาจใช้เวลาสักครู่ถ้าไฟล์ใหญ่)")
    wb = load_workbook(path, read_only=True, data_only=True)
    found = {t: [] for t in targets}
    for ws in wb.worksheets:
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            for c_idx, value in enumerate(row, start=1):
                if value is None:
                    continue
                d = _digits(value)
                for t in targets:
                    if d == t or (len(t) >= 8 and t in d):
                        if len(found[t]) < 20:
                            found[t].append((ws.title, r_idx, get_column_letter(c_idx), value))
    for t, hits in found.items():
        if not hits:
            print(f"    ✗ ไม่พบ {t} ในไฟล์นี้")
        for sheet, r, col, value in hits:
            print(f"    ✓ พบ {t}  →  ชีต '{sheet}' ช่อง {col}{r}  (ค่าในช่อง: {value!r})")

    # แสดงทั้งแถวของรายการแรกที่เจอ พร้อมชื่อหัวคอลัมน์ เพื่อเทียบกับที่แดชบอร์ดแสดง
    first = next((h[0] for h in found.values() if h), None)
    if first:
        sheet, r, _, _ = first
        ws = wb[sheet]
        header, header_row = None, None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
            names = [re.sub(r"[\s_\-]+", "", str(v)).casefold() for v in row if v is not None]
            if "totalcost" in names or "loading" in names:
                header, header_row = row, i
                break
        values = next(ws.iter_rows(min_row=r, max_row=r, values_only=True))
        print(f"\n    ทั้งแถวที่ {r} ในชีต '{sheet}' (หัวคอลัมน์จากแถวที่ {header_row}):")
        for c_idx, value in enumerate(values, start=1):
            if value is None or str(value).strip() == "":
                continue
            name = header[c_idx - 1] if header and c_idx - 1 < len(header) else "(ไม่มีหัวคอลัมน์)"
            mark = "   ◀ Total Cost" if re.sub(r"[\s_\-]+", "", str(name)).casefold() == "totalcost" else ""
            print(f"      {get_column_letter(c_idx):>3} | {name}: {value!r}{mark}")
    wb.close()


def search_database(targets):
    print(f"\n[2] ค้นในฐานข้อมูล: {DB_PATH}")
    if not DB_PATH.exists():
        print("    ✗ ยังไม่มีไฟล์ฐานข้อมูล")
        return
    import duckdb
    import pandas as pd

    con = duckdb.connect(str(DB_PATH), read_only=True)
    try:
        df = con.execute("SELECT * FROM trips").df()
    finally:
        con.close()

    df = df[df["_source_file"].map(_key).str.contains(_key(SOURCE_KEYWORD), na=False)]
    print(f"    ข้อมูลจากไฟล์กล้วยไม้ในฐานข้อมูล: {len(df):,} แถว")
    if "_source_row" not in df.columns:
        print("    ⚠ ฐานข้อมูลยังไม่มีเลขแถว (_source_row) — กด 'สร้างฐานข้อมูลใหม่' ในหน้า Data Management ก่อน")

    as_digits = df.astype("string").apply(lambda col: col.map(_digits))
    for t in targets:
        mask = (as_digits == t).any(axis=1)
        hits = df[mask]
        if hits.empty:
            print(f"    ✗ ไม่พบ {t} ในฐานข้อมูล")
            continue
        print(f"    ✓ พบ {t} ในฐานข้อมูล {len(hits):,} แถว (แสดงสูงสุด 3 แถว)")
        for _, row in hits.head(3).iterrows():
            print("      " + "-" * 60)
            for col, value in row.items():
                if pd.notna(value) and str(value).strip() not in {"", "nan", "None"}:
                    print(f"      {col}: {value}")


if __name__ == "__main__":
    targets = [_digits(a) for a in sys.argv[1:]] or DEFAULT_TARGETS
    print("ค่าที่ค้นหา:", ", ".join(targets))

    source = find_source_file()
    if source is None:
        print(f"✗ ไม่พบไฟล์ที่ชื่อมีคำว่า '{SOURCE_KEYWORD}' ในโฟลเดอร์ {DATA_FOLDER.resolve()}")
    else:
        search_excel(source, targets)
    search_database(targets)
    print("\nเสร็จแล้ว — แคปผลลัพธ์ทั้งหมดนี้ส่งกลับมาได้เลย")