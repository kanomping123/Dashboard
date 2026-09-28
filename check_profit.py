"""
check_profit.py
---------------
ตรวจหาสาเหตุที่หน้า "กำไรลูกค้าและประเภทสินค้า" ค้าง (รันในโฟลเดอร์เดียวกับ app.py)

    python check_profit.py

จะบอก: ขนาดไฟล์, ชีต/จำนวนแถว, เวลาที่ใช้อ่านไฟล์, คอลัมน์ที่จับได้ และเวลาที่ใช้คำนวณ
"""

import time
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

print("=" * 60)
print("1) หาไฟล์ในโฟลเดอร์ data")
data = Path("data")
files = [f for f in data.glob("*.xls*") if not f.name.startswith("~$")] if data.exists() else []
for f in files:
    print(f"   - {f.name}  ({f.stat().st_size / 1024 / 1024:,.1f} MB)")

import customer_profit as cp  # noqa: E402

path = cp.find_source_file()
if path is None:
    print("   ✗ ไม่พบไฟล์ที่ชื่อมีคำว่า 'กำไรลูกค้า'")
    raise SystemExit
print(f"   ✓ ใช้ไฟล์: {path.name}")

print("=" * 60)
print("2) ชีตและจำนวนแถวในไฟล์ (อ่านแบบเร็ว)")
from openpyxl import load_workbook  # noqa: E402

t = time.time()
wb = load_workbook(path, read_only=True, data_only=True)
for ws in wb.worksheets:
    print(f"   - ชีต '{ws.title}': ประมาณ {ws.max_row:,} แถว × {ws.max_column:,} คอลัมน์")
wb.close()
print(f"   (ใช้เวลา {time.time() - t:,.1f} วินาที)")

print("=" * 60)
print("3) อ่านไฟล์แบบเดียวกับแดชบอร์ด")
t = time.time()
loader = getattr(cp.load_customer_profit, "__wrapped__", cp.load_customer_profit)
df, used, sheet = loader(str(path), path.stat().st_mtime_ns)
print(f"   ใช้เวลา {time.time() - t:,.1f} วินาที · ชีต '{sheet}' · {len(df):,} ลูกค้า")
for k, v in used.items():
    print(f"   - {k:<9} ← คอลัมน์ '{v}'")

if df.empty:
    print("   ✗ อ่านข้อมูลไม่ได้ (หาแถวหัวตารางไม่เจอ)")
    raise SystemExit

print("=" * 60)
print("4) ข้อมูลตัวอย่าง")
print(df.head(5).to_string())
print(f"\n   ประเภทสินค้า: {df['Type'].nunique():,} ประเภท")
print(f"   รายได้รวม {df['Revenue'].sum():,.0f} · ต้นทุน {df['Cost'].sum():,.0f} · กำไร {df['Profit'].sum():,.0f}")

print("=" * 60)
print("เสร็จแล้ว — แคปผลลัพธ์ทั้งหมดนี้ส่งกลับมาได้เลย")