"""
ar_dashboard.py
---------------
แดชบอร์ด "ลูกหนี้ & การเก็บเงิน" (AR & Collection) — อ่านไฟล์รายงานลูกหนี้ตรงจากโฟลเดอร์ data

แหล่งข้อมูล
- ไฟล์ Excel ในโฟลเดอร์ data ที่ชื่อมีคำว่า "ลูกหนี้" (หลายไฟล์ = ใช้ไฟล์ที่แก้ไขล่าสุด)
  ถ้าไม่เจอจากชื่อไฟล์ จะหาไฟล์ที่มีชีตชื่อ "รายงานลูกหนี้"
- ชีต "รายงานลูกหนี้"  : 1 แถว = 1 บิล
    ชื่อลูกหนี้ · สาขา · ระยะเวลาเครดิต · เลขที่ใบวางบิล · วันที่วางบิลได้ · วันที่จบ(วันที่จ่าย)
    จำนวนเงิน · วันครบกำหนด(Due Date) · วันที่ปัจจุบัน · กลุ่มอายุลูกหนี้ · ตัวเลขวันช้า · ชั้นลูกหนี้ (TFRS 9)
- ชีต "สรุปลูกหนี้"     : 1 แถว = 1 ลูกหนี้
    จำนวนบิลทั้งหมด · จำนวนบิลที่ช้า · รวมวันช้า · วันช้าสูงสุด · เปอร์เซนต์บิลที่ช้า
    กลุ่มช่วงอายุลูกหนี้ · การแบ่งชั้นลูกหนี้ (TFRS 9)
- ชีต "เงื่อนไข"        : คำอธิบายช่วงอายุหนี้และชั้นลูกหนี้ (แสดงในหน้าเป็นเกณฑ์อ้างอิง)

ตรรกะ
- ยังไม่ชำระ   = ไม่มี วันที่จบ(วันที่จ่าย)
- วันช้า        = คอลัมน์ ตัวเลขวันช้า (ไม่มีค่า → คำนวณจาก วันที่จ่าย/วันที่ปัจจุบัน − วันครบกำหนด)
- เกินกำหนด    = ยังไม่ชำระ และ วันช้า > 0
- ช่วงอายุหนี้ = จัดจากวันช้าตามเกณฑ์ในชีตเงื่อนไข (0 / 1–30 / 31–60 / 61–90 / เกิน 90 วัน)
- ชั้นลูกหนี้   = ระดับบิล: คอลัมน์ ชั้นลูกหนี้ (TFRS 9)
                  ระดับลูกหนี้: การแบ่งชั้นลูกหนี้ ในชีตสรุป (ไม่มี → ใช้ชั้นที่แย่ที่สุดของบิล)
- ชื่อคอลัมน์จับแบบยืดหยุ่น (เว้นวรรค/วงเล็บ/สะกดต่างเล็กน้อยได้) เช่น วันที่จบ = วันที่จ่าย
- แถวในชีตสรุปที่มีแต่ชื่อ ไม่มีตัวเลข จะไม่ถูกนับ
- แต่ละส่วนของหน้าแยกกันทำงาน: ถ้าส่วนใดผิดพลาด จะขึ้นเตือนเฉพาะส่วนนั้น หน้าที่เหลือยังแสดงได้

ส่วนสำหรับผู้บริหาร
- สรุปอัตโนมัติ: ยอดค้าง/เกินกำหนด, NPL, การกระจุกตัว, ประสิทธิภาพการเก็บเงิน, เทียบเดือนก่อน
- ต้องลงมือ: แบ่งลูกหนี้ตามความเร่งด่วน (เกิน 90 วัน / 31–90 วัน / 1–30 วัน / จ่ายช้าเป็นประจำ)
- การกระจุกตัวของยอดคงค้าง (Pareto)
"""

import re
import unicodedata
from html import escape as esc
from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import streamlit as st

from transport_cost_route import (
    PAGE_CSS,
    THAI_MONTHS,
    _card,
    _money,
    _money_full,
    _num,
    _style,
    _table_html,
    _tint,
    _wkey,
)

DATA_FOLDER = Path("data")
FILE_KEYWORDS = ["ลูกหนี้"]
PLOT_CONFIG = {"displayModeBar": False}

# ---------------- ช่วงอายุหนี้ (ตามชีตเงื่อนไข) ----------------
AGING_KEYS = ["ชำระตรงเวลา", "ค้างชำระ 1 - 30 วัน", "ค้างชำระ 31 - 60 วัน",
              "ค้างชำระ 61 - 90 วัน", "ค้างชำระเกิน 90 วัน"]
AGING_LABELS = ["ยังไม่ถึงกำหนด", "1–30 วัน", "31–60 วัน", "61–90 วัน", "เกิน 90 วัน"]
AGING_COLORS = ["#86CFA3", "#EFCB64", "#F6AE6B", "#EE8A7C", "#C23B53"]

# ---------------- ชั้นลูกหนี้ (TFRS 9) ----------------
CLASS_NAMES = ["ลูกหนี้ชั้นดี", "ลูกหนี้เฝ้าติดตาม", "ลูกหนี้ด้อยคุณภาพ (NPL)", "ไม่ระบุ"]
CLASS_COLORS = ["#86CFA3", "#F6AE6B", "#E0566C", "#CBD5E1"]
CLASS_TEXT = ["#1F7A4D", "#B45309", "#B71C1C", "#64748B"]
CLASS_DEFAULT_RULE = ["ไม่ค้างชำระ หรือ ค้างไม่เกิน 30 วัน", "ค้างชำระ 31 - 90 วัน", "ค้างชำระเกิน 90 วัน", ""]

STATUS_COLORS = {"ยังไม่ถึงกำหนด": "#86CFA3", "เกินกำหนด": "#E0566C",
                 "ชำระตรงเวลา": "#7CC4D6", "ชำระล่าช้า": "#F6AE6B"}

AR_CSS = """
<style>
.ar-badge { display: inline-block; font-size: 12px; font-weight: 700; padding: 2px 10px; border-radius: 99px; white-space: nowrap; }
.ar-name { max-width: 300px; overflow: hidden; text-overflow: ellipsis; display: inline-block; vertical-align: bottom; }
.ar-over { color: #C23B53 !important; font-weight: 700; }
.ar-leg-row { display: grid; grid-template-columns: 12px 1fr auto; gap: 10px; align-items: start;
              padding: 8px 4px; border-bottom: 1px dashed #EEF2F7; font-size: 13px; }
.ar-leg-row .sw { width: 12px; height: 12px; border-radius: 4px; margin-top: 4px; }
.ar-leg-row .nm { color: #1E293B; font-weight: 600; }
.ar-leg-row .rule { color: #94A3B8; font-size: 11.5px; font-weight: 400; line-height: 1.45; }
.ar-leg-row .val { text-align: right; color: #334155; font-variant-numeric: tabular-nums; }
.ar-leg-row .val small { display: block; color: #94A3B8; font-size: 11.5px; }
.ar-crit { width: 100%; border-collapse: collapse; font-size: 13px; }
.ar-crit td, .ar-crit th { padding: 8px 10px; border-bottom: 1px solid #FBEFF1; text-align: left; vertical-align: top; }
.ar-crit th { background: #FFF3F5; color: #475569; font-weight: 700; }
/* สรุปสำหรับผู้บริหาร */
.ar-exec { background: #fff; border: 1px solid #F4D8DE; border-radius: 18px; padding: 16px 20px; margin: 4px 0 14px;
           box-shadow: 0 1px 2px rgba(15,23,42,.04), 0 8px 24px rgba(226,90,112,.08); }
.ar-exec-head { font-size: 17px; font-weight: 800; color: #0F172A; margin-bottom: 6px; }
.ar-exec-row { display: grid; grid-template-columns: 30px 1fr; gap: 10px; align-items: start; padding: 9px 12px;
               border-radius: 12px; margin-top: 6px; font-size: 14px; color: #1E293B; line-height: 1.55; }
.ar-exec-row .ic { font-size: 18px; line-height: 1.3; }
.ar-exec-row.good { background: #F1FAF4; } .ar-exec-row.warn { background: #FFF8EC; } .ar-exec-row.bad { background: #FFF1F2; }
.ar-exec-row.info { background: #F8FAFC; }
.ar-exec-row b { font-variant-numeric: tabular-nums; }
.ar-up { color: #C23B53; font-weight: 700; } .ar-down { color: #2E8B57; font-weight: 700; }
/* ต้องลงมือ */
.ar-act-grid { display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: 12px; margin-top: 8px; }
.ar-act { border: 1px solid #F1E4E7; border-radius: 16px; padding: 12px 14px; background: #fff; min-width: 0;
          border-top: 4px solid var(--c); }
.ar-act-title { font-size: 14px; font-weight: 800; color: #0F172A; }
.ar-act-do { font-size: 12px; color: #64748B; margin: 2px 0 8px; line-height: 1.45; }
.ar-act-num { font-size: 22px; font-weight: 800; color: var(--c); font-variant-numeric: tabular-nums; }
.ar-act-num span { font-size: 12.5px; font-weight: 600; color: #64748B; }
.ar-act-list { margin-top: 8px; border-top: 1px dashed #F1E4E7; }
.ar-act-item { display: grid; grid-template-columns: 1fr auto; gap: 8px; font-size: 12.5px; padding: 5px 0;
               border-bottom: 1px dashed #F6EEF0; }
.ar-act-item .n { overflow: hidden; text-overflow: ellipsis; white-space: nowrap; color: #1E293B; }
.ar-act-item .v { text-align: right; color: #334155; font-variant-numeric: tabular-nums; white-space: nowrap; }
.ar-act-item .v small { display: block; color: #94A3B8; font-size: 11px; }
.ar-act-more { font-size: 11.5px; color: #94A3B8; margin-top: 6px; }
.ar-act-empty { font-size: 12.5px; color: #2E8B57; margin-top: 8px; }
@media (max-width: 1100px) { .ar-act-grid { grid-template-columns: repeat(2, minmax(0, 1fr)); } }
@media (max-width: 640px) { .ar-act-grid { grid-template-columns: 1fr; } }
</style>
"""


# =========================================================
# HELPERS
# =========================================================

def _n(value) -> str:
    return re.sub(r"[\s_\-\.\(\)]+", "", unicodedata.normalize("NFC", str(value))).casefold()


_RATE_TOKENS = ("%", "เปอร์เซ", "สัดส่วน", "rate", "percent")

DETAIL_SPEC = [
    ("debtor", ["ชื่อลูกหนี้", "ลูกหนี้", "ชื่อลูกค้า", "ลูกค้า"], lambda k: "ลูกหนี้" in k and "ชื่อ" in k),
    ("branch", ["สาขา", "branch"], lambda k: "สาขา" in k),
    ("credit", ["ระยะเวลาเครดิต", "เครดิต", "creditterm"], lambda k: "เครดิต" in k),
    ("bill", ["เลขที่ใบวางบิล", "เลขที่บิล", "เลขที่ใบแจ้งหนี้"], lambda k: "เลขที่" in k),
    ("inv_date", ["วันที่วางบิลได้", "วันที่วางบิล"], lambda k: "วางบิล" in k and "วันที่" in k),
    ("pay_date", ["วันที่จบวันที่จ่าย", "วันที่จ่าย", "วันที่ชำระ", "วันที่จบ", "วันที่รับชำระ"],
     lambda k: "จ่าย" in k or "จบ" in k or ("ชำระ" in k and "วันที่" in k)),
    ("amount", ["จำนวนเงิน", "ยอดเงิน", "amount"], lambda k: "จำนวนเงิน" in k or "ยอดเงิน" in k),
    ("due_date", ["วันครบกำหนดDueDate", "วันครบกำหนด", "duedate"], lambda k: "ครบกำหนด" in k or "due" in k),
    ("ref_date", ["วันที่ปัจจุบัน"], lambda k: "ปัจจุบัน" in k),
    ("aging", ["กลุ่มอายุลูกหนี้"], lambda k: "อายุ" in k),
    ("days", ["ตัวเลขวันช้า", "จำนวนวันช้า", "วันช้า"], lambda k: "ช้า" in k),
    ("cls", ["ชั้นลูกหนี้TFRS9", "ชั้นลูกหนี้", "แบ่งชั้นลูกหนี้"], lambda k: "ชั้น" in k),
]

SUMMARY_SPEC = [
    ("debtor", ["ชื่อลูกหนี้", "ลูกหนี้", "ชื่อลูกค้า", "ลูกค้า"], lambda k: "ลูกหนี้" in k and "ชื่อ" in k),
    ("late_rate", ["เปอร์เซนต์บิลที่ช้า", "เปอร์เซ็นต์บิลที่ช้า", "%บิลที่ช้า"],
     lambda k: "ช้า" in k and any(t in k for t in _RATE_TOKENS)),
    ("late_bills", ["จำนวนบิลที่ช้า", "บิลที่ช้า"], lambda k: "บิล" in k and "ช้า" in k),
    ("bills", ["จำนวนบิลทั้งหมด", "จำนวนบิล"], lambda k: "บิล" in k and "ช้า" not in k),
    ("max_delay", ["วันช้าสูงสุด"], lambda k: "ช้า" in k and "สูงสุด" in k),
    ("total_delay", ["รวมวันช้า"], lambda k: "ช้า" in k and "รวม" in k),
    ("aging", ["กลุ่มช่วงอายุลูกหนี้", "กลุ่มอายุลูกหนี้"], lambda k: "อายุ" in k),
    ("cls", ["การแบ่งชั้นลูกหนี้TFRS9", "แบ่งชั้นลูกหนี้", "ชั้นลูกหนี้"], lambda k: "ชั้น" in k),
]


def _map_columns(df: pd.DataFrame, spec) -> dict:
    """ชื่อตรงก่อน แล้วค่อยเดาจากคำในชื่อ (1 คอลัมน์ใช้ได้ครั้งเดียว)"""
    norm = {c: _n(c) for c in df.columns if not str(c).startswith("Unnamed")}
    found, used = {}, set()
    for key, names, _rule in spec:
        targets = {_n(x) for x in names}
        hit = next((c for c, k in norm.items() if k in targets and c not in used), None)
        if hit is not None:
            found[key] = hit
            used.add(hit)
    for key, _names, rule in spec:
        if key in found:
            continue
        hit = next((c for c, k in norm.items() if c not in used and rule(k)), None)
        if hit is not None:
            found[key] = hit
            used.add(hit)
    return found


def _to_date(series: pd.Series) -> pd.Series:
    """วันที่จาก Excel: datetime / ข้อความ dd/mm/yyyy / เลข serial / ปี พ.ศ."""
    if pd.api.types.is_datetime64_any_dtype(series):
        dt = pd.to_datetime(series, errors="coerce")
    else:
        num = pd.to_numeric(series, errors="coerce")
        serial = num.where(num.between(20000, 80000))
        dt = pd.to_datetime(series.where(serial.isna()), errors="coerce", dayfirst=True)
        if serial.notna().any():
            dt = dt.fillna(pd.to_datetime(serial, unit="D", origin="1899-12-30", errors="coerce"))
    be = dt.dt.year > 2400
    if be.any():
        dt = dt.where(~be, dt - pd.DateOffset(years=543))
    return dt


def _class_idx(value) -> int:
    """ข้อความชั้นลูกหนี้ -> 0 ชั้นดี / 1 เฝ้าติดตาม / 2 ด้อยคุณภาพ / 3 ไม่ระบุ"""
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return 3
    k = _n(value)
    if not k:
        return 3
    if any(t in k for t in ("ด้อย", "แย่", "npl", "nonperforming", "สูญ")):
        return 2
    if any(t in k for t in ("เฝ้า", "ติดตาม", "underperforming")):
        return 1
    if "ดี" in k or "performing" in k:
        return 0
    return 3


def _class_from_days(days) -> int:
    if pd.isna(days):
        return 3
    return 0 if days <= 30 else (1 if days <= 90 else 2)


def _aging_idx_from_days(days) -> int:
    if pd.isna(days) or days <= 0:
        return 0
    if days <= 30:
        return 1
    if days <= 60:
        return 2
    if days <= 90:
        return 3
    return 4


def _aging_idx_from_text(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    k = _n(value)
    if not k:
        return None
    if "ตรงเวลา" in k or "ยังไม่ถึง" in k:
        return 0
    if "เกิน90" in k or "90+" in k or "91" in k:
        return 4
    if "61" in k:
        return 3
    if "31" in k:
        return 2
    if "130" in k or "1-30" in k:
        return 1
    return None


def _id_text(value) -> str:
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def _badge(text, color, fg=None) -> str:
    return (f'<span class="ar-badge" style="background:{_tint(color, 0.75)};'
            f'color:{fg or "#3F2A2E"}">{esc(str(text))}</span>')


def _class_badge(idx) -> str:
    idx = int(idx) if pd.notna(idx) else 3
    return _badge(CLASS_NAMES[idx], CLASS_COLORS[idx], CLASS_TEXT[idx])


def _name_html(name) -> str:
    s = str(name)
    return f'<span class="ar-name" title="{esc(s)}">{esc(s)}</span>'


class _safe:
    """ครอบแต่ละส่วนของหน้า: ถ้าส่วนนั้นผิดพลาด ขึ้นเตือนเฉพาะส่วน แล้วแสดงส่วนอื่นต่อ"""

    def __init__(self, name):
        self.name = name

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, _tb):
        if exc_type is None:
            return False
        # ปล่อยคำสั่งภายในของ Streamlit (rerun/stop) และการหยุดโปรแกรมผ่านไปตามปกติ
        if (issubclass(exc_type, (KeyboardInterrupt, SystemExit))
                or getattr(exc_type, "__module__", "").startswith("streamlit")):
            return False
        st.warning(f"ส่วน “{self.name}” แสดงไม่ได้ ({exc_type.__name__}: {exc}) — ส่วนอื่นยังใช้งานได้ตามปกติ")
        return True


def _pct(v, digits=0) -> str:
    return "—" if v is None or pd.isna(v) else f"{v * 100:.{digits}f}%"


def _wavg(values: pd.Series, weights: pd.Series) -> float:
    m = values.notna() & weights.notna() & (weights > 0)
    w = weights[m].sum()
    return float((values[m] * weights[m]).sum() / w) if w else float("nan")


def _thai_date(ts) -> str:
    if ts is None or pd.isna(ts):
        return "—"
    return f"{ts.day} {THAI_MONTHS[ts.month]} {ts.year + 543 if ts.year < 2400 else ts.year}"


# =========================================================
# LOAD
# =========================================================

def find_ar_file():
    """ไฟล์ลูกหนี้ในโฟลเดอร์ data (ชื่อมีคำว่า ลูกหนี้ ก่อน · ไม่เจอค่อยดูชื่อชีต)"""
    if not DATA_FOLDER.exists():
        return None
    files = [
        f for f in DATA_FOLDER.iterdir()
        if f.is_file() and f.suffix.lower() in {".xlsx", ".xlsm"} and not f.name.startswith("~$")
    ]
    named = [f for f in files if any(_n(k) in _n(f.stem) for k in FILE_KEYWORDS)]
    if named:
        return max(named, key=lambda f: f.stat().st_mtime_ns)
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    hits = []
    for f in files:
        try:
            wb = load_workbook(f, read_only=True)
            if any("รายงานลูกหนี้" in _n(s) for s in wb.sheetnames):
                hits.append(f)
            wb.close()
        except Exception:
            continue
    return max(hits, key=lambda f: f.stat().st_mtime_ns) if hits else None


def _read_sheet(xls, sheet, must_have="ลูกหนี้"):
    """อ่านชีต โดยหาแถวหัวตารางอัตโนมัติ (แถวแรกที่มีคำว่า ลูกหนี้)"""
    raw = pd.read_excel(xls, sheet_name=sheet, header=None, nrows=15)
    header_row = 0
    for i, row in raw.iterrows():
        if any(must_have in _n(v) for v in row.tolist() if isinstance(v, str)):
            header_row = i
            break
    df = pd.read_excel(xls, sheet_name=sheet, header=header_row)
    df.columns = [" ".join(str(c).split()) for c in df.columns]
    return df.dropna(how="all")


@st.cache_data(show_spinner=False)
def load_ar_file(path_str: str, mtime_ns: int):
    xls = pd.ExcelFile(path_str)
    sheets = xls.sheet_names
    detail_sheet = next((s for s in sheets if "รายงาน" in _n(s) and "ลูกหนี้" in _n(s)), None)
    if detail_sheet is None:  # ชีตที่มีคอลัมน์ จำนวนเงิน + ครบกำหนด
        for s in sheets:
            cols = [_n(c) for c in pd.read_excel(xls, sheet_name=s, nrows=0).columns]
            if any("จำนวนเงิน" in c for c in cols) and any("ครบกำหนด" in c for c in cols):
                detail_sheet = s
                break
    summary_sheet = next((s for s in sheets if "สรุป" in _n(s)), None)
    cond_sheet = next((s for s in sheets if "เงื่อนไข" in _n(s)), None)

    detail = _read_sheet(xls, detail_sheet) if detail_sheet else pd.DataFrame()
    summary = _read_sheet(xls, summary_sheet) if summary_sheet else pd.DataFrame()
    cond = pd.read_excel(xls, sheet_name=cond_sheet, header=None) if cond_sheet else pd.DataFrame()
    return detail, summary, cond, detail_sheet, summary_sheet


def parse_conditions(cond: pd.DataFrame):
    """ชีตเงื่อนไข -> (เกณฑ์ช่วงอายุ, เกณฑ์ชั้นลูกหนี้) เป็น list ของ (ชื่อ, เกณฑ์, ความหมาย)"""
    aging_rows, class_rows = [], []
    if cond is None or cond.empty or cond.shape[1] < 3:
        return aging_rows, class_rows
    for _, r in cond.iterrows():
        a, b, c = (r.iloc[0], r.iloc[1], r.iloc[2])
        if not all(isinstance(x, str) and x.strip() for x in (a, b, c)):
            continue
        if _n(a) in {_n("ชั้นลูกหนี้")}:  # แถวหัวตาราง
            continue
        if "ลูกหนี้" in a:
            class_rows.append((a.strip(), b.strip(), c.strip()))
        else:
            aging_rows.append((a.strip(), b.strip(), c.strip()))
    return aging_rows, class_rows


def prepare(detail: pd.DataFrame, summary: pd.DataFrame):
    dc = _map_columns(detail, DETAIL_SPEC)
    missing = [label for key, label in (("debtor", "ชื่อลูกหนี้"), ("amount", "จำนวนเงิน"),
                                         ("due_date", "วันครบกำหนด")) if key not in dc]
    if missing:
        return None, None, dc, missing

    ar = pd.DataFrame(index=detail.index)
    ar["Debtor"] = detail[dc["debtor"]].astype("string").fillna("").str.strip()
    ar = ar[ar["Debtor"] != ""]
    d = detail.loc[ar.index]

    def text(key, default="ไม่ระบุ"):
        if key not in dc:
            return pd.Series(default, index=ar.index, dtype="string")
        s = d[dc[key]].astype("string").fillna("").str.strip()
        return s.mask(s.eq(""), default)

    ar["Branch"] = text("branch")
    ar["Credit"] = _num(d[dc["credit"]]) if "credit" in dc else float("nan")
    ar["Bill"] = d[dc["bill"]].map(_id_text) if "bill" in dc else ""
    ar["Amount"] = _num(d[dc["amount"]]).fillna(0.0)
    for key, col in (("inv_date", "InvDate"), ("pay_date", "PayDate"),
                     ("due_date", "DueDate"), ("ref_date", "RefDate")):
        ar[col] = _to_date(d[dc[key]]) if key in dc else pd.NaT

    ref_default = ar["RefDate"].dropna().max()
    if pd.isna(ref_default):
        others = pd.concat([ar["InvDate"], ar["PayDate"], ar["DueDate"]]).dropna()
        ref_default = others.max() if not others.empty else pd.Timestamp.today().normalize()
    ar["RefDate"] = ar["RefDate"].fillna(ref_default)

    ar["Paid"] = ar["PayDate"].notna()
    calc_days = (ar["PayDate"].where(ar["Paid"], ar["RefDate"]) - ar["DueDate"]).dt.days
    file_days = _num(d[dc["days"]]) if "days" in dc else pd.Series(float("nan"), index=ar.index)
    ar["DaysLate"] = file_days.fillna(calc_days).fillna(0).clip(lower=0)

    ar["Unpaid"] = ~ar["Paid"]
    ar["IsLate"] = ar["DaysLate"] > 0
    ar["Outstanding"] = ar["Amount"].where(ar["Unpaid"], 0.0)
    ar["OverdueFlag"] = ar["Unpaid"] & ar["IsLate"]
    ar["OverdueAmt"] = ar["Amount"].where(ar["OverdueFlag"], 0.0)
    ar["AgeIdx"] = ar["DaysLate"].map(_aging_idx_from_days).where(ar["Unpaid"], -1).astype(int)
    ar["AgingText"] = text("aging", "")
    ar["Status"] = "ยังไม่ถึงกำหนด"
    ar.loc[ar["OverdueFlag"], "Status"] = "เกินกำหนด"
    ar.loc[ar["Paid"] & ~ar["IsLate"], "Status"] = "ชำระตรงเวลา"
    ar.loc[ar["Paid"] & ar["IsLate"], "Status"] = "ชำระล่าช้า"

    bill_cls = d[dc["cls"]].map(_class_idx) if "cls" in dc else pd.Series(3, index=ar.index)
    ar["ClassIdx"] = bill_cls.where(bill_cls != 3, ar["DaysLate"].map(_class_from_days)).astype(int)

    # ---------------- ระดับลูกหนี้ ----------------
    sc = _map_columns(summary, SUMMARY_SPEC) if summary is not None and not summary.empty else {}
    src = None
    if "debtor" in sc:
        src = pd.DataFrame({"Debtor": summary[sc["debtor"]].astype("string").fillna("").str.strip()})
        for key, col in (("bills", "SrcBills"), ("late_bills", "SrcLateBills"),
                         ("total_delay", "SrcTotalDelay"), ("max_delay", "SrcMaxDelay"),
                         ("late_rate", "SrcLateRate")):
            if key in sc:
                src[col] = _num(summary[sc[key]].astype("string").str.replace("%", "", regex=False))
        if "SrcLateRate" in src.columns:
            rate = src["SrcLateRate"].dropna()
            if not rate.empty and rate.abs().median() > 1.5:  # เก็บเป็น 25 แทน 0.25
                src["SrcLateRate"] = src["SrcLateRate"] / 100
        if "aging" in sc:
            src["SrcAging"] = summary[sc["aging"]].astype("string").fillna("").str.strip()
        if "cls" in sc:
            src["SrcClass"] = summary[sc["cls"]].map(_class_idx)
        src = src[src["Debtor"] != ""]
        # แถวที่มีแต่ชื่อ ไม่มีตัวเลข/กลุ่ม/ชั้นเลย = แถวว่างท้ายชีต ไม่นับ
        has_info = pd.Series(False, index=src.index)
        for col in src.columns:
            if col.startswith("SrcAging"):
                has_info |= src[col].fillna("").ne("")
            elif col == "SrcClass":
                has_info |= src[col].ne(3)
            elif col.startswith("Src"):
                has_info |= src[col].notna()
        src = src[has_info]
        src = src.drop_duplicates("Debtor", keep="last")

    worst = ar.groupby("Debtor")["ClassIdx"].agg(lambda s: max((v for v in s if v != 3), default=3))
    debtor_class = worst.copy()
    if src is not None and "SrcClass" in src.columns:
        from_src = src.set_index("Debtor")["SrcClass"]
        from_src = from_src[from_src != 3]
        debtor_class.update(from_src.reindex(debtor_class.index).dropna().astype(int))
    ar["DebtorClass"] = ar["Debtor"].map(debtor_class).fillna(3).astype(int)
    return ar, src, dc, []


def debtor_table(f: pd.DataFrame, src) -> pd.DataFrame:
    """สรุปรายลูกหนี้: ยอดค้าง/เกินกำหนดจากข้อมูลตามตัวกรอง · สถิติบิลจากชีตสรุป (ไม่มี → คำนวณ)"""
    g = (
        f.groupby("Debtor")
        .agg(
            Branch=("Branch", lambda s: ", ".join(sorted(set(s))[:3])),
            Bills=("Amount", "size"), LateBills=("IsLate", "sum"),
            TotalDelay=("DaysLate", "sum"), MaxDelay=("DaysLate", "max"),
            Outstanding=("Outstanding", "sum"), OverdueAmt=("OverdueAmt", "sum"),
            UnpaidBills=("Unpaid", "sum"), Class=("DebtorClass", "first"),
            MaxUnpaidDelay=("DaysLate", lambda s: float(s[f.loc[s.index, "Unpaid"]].max())
                            if f.loc[s.index, "Unpaid"].any() else 0.0),
        )
        .reset_index()
    )
    g["LateRate"] = g["LateBills"] / g["Bills"].replace(0, float("nan"))
    g["AgingIdx"] = g["MaxUnpaidDelay"].map(_aging_idx_from_days)
    g["AgingText"] = g["AgingIdx"].map(lambda i: AGING_KEYS[i])
    if src is not None:
        g = g.merge(src, on="Debtor", how="left")
        for base, col in (("Bills", "SrcBills"), ("LateBills", "SrcLateBills"),
                          ("TotalDelay", "SrcTotalDelay"), ("MaxDelay", "SrcMaxDelay"),
                          ("LateRate", "SrcLateRate")):
            if col in g.columns:
                g[base] = g[col].fillna(g[base])
        if "SrcAging" in g.columns:
            g["AgingText"] = g["SrcAging"].mask(g["SrcAging"].fillna("").eq(""), g["AgingText"])
    return g


# =========================================================
# DASHBOARD
# =========================================================

def render_ar_dashboard():
    path = find_ar_file()
    if path is None:
        st.error("ไม่พบไฟล์ลูกหนี้ในโฟลเดอร์ data — อัปโหลดไฟล์ที่ชื่อมีคำว่า “ลูกหนี้” "
                 "(เช่น รายงานลูกหนี้.xlsx) ในหน้า Data Management ก่อน")
        return
    try:
        detail, summary, cond, detail_sheet, summary_sheet = load_ar_file(str(path), path.stat().st_mtime_ns)
    except Exception as e:
        st.error(f"อ่านไฟล์ {path.name} ไม่ได้: {e}")
        return
    if detail.empty:
        st.error(f"ไฟล์ {path.name} ไม่มีชีตรายงานลูกหนี้ (รายละเอียดรายบิล)")
        return

    ar, src, dc, missing = prepare(detail, summary)
    if missing:
        st.error(f"ชีต {detail_sheet} ขาดคอลัมน์: " + ", ".join(missing)
                 + " · คอลัมน์ที่พบ: " + ", ".join(map(str, detail.columns)))
        return
    if ar.empty:
        st.warning("ไม่มีรายการลูกหนี้ในไฟล์")
        return
    aging_crit, class_crit = parse_conditions(cond)
    ref_date = ar["RefDate"].max()

    st.markdown(PAGE_CSS + AR_CSS, unsafe_allow_html=True)
    st.caption(f"ข้อมูลลูกหนี้ ณ วันที่ {_thai_date(ref_date)} · ไฟล์ {path.name}")

    # ---------------- ตัวกรอง ----------------
    fcard = _card("ar_filters")
    fcard.markdown('<div class="tc-filter-title">🔎 ตัวกรองข้อมูล</div>', unsafe_allow_html=True)
    f1, f2, f3, f4, f5, f6 = fcard.columns([0.8, 1.1, 1.2, 1.1, 1.4, 1.6])
    filtered = ar
    inv_year = ar["InvDate"].dt.year
    ar_year = inv_year.map(lambda y: "" if pd.isna(y) else str(int(y + 543)))
    with f1:
        years = sorted(y for y in ar_year.unique() if y)
        year = st.selectbox("ปีวางบิล", ["ทั้งหมด"] + years, key="ar2_year")
    if year != "ทั้งหมด":
        filtered = filtered[ar_year.loc[filtered.index] == year]
    with f2:
        months = sorted(int(m) for m in filtered["InvDate"].dt.month.dropna().unique())
        chosen_m = st.multiselect("เดือนวางบิล", [THAI_MONTHS[m] for m in months],
                                  placeholder="ทั้งหมด", key="ar2_month")
    if chosen_m:
        nums = [n for n, lb in THAI_MONTHS.items() if lb in chosen_m]
        filtered = filtered[filtered["InvDate"].dt.month.isin(nums)]
    with f3:
        branches = sorted(b for b in filtered["Branch"].unique() if b)
        chosen_b = st.multiselect("สาขา", branches, placeholder="ทั้งหมด", key=_wkey("ar2_branch", branches))
    if chosen_b:
        filtered = filtered[filtered["Branch"].isin(chosen_b)]
    with f4:
        status_opts = ["ทั้งหมด", "ยังไม่ชำระ", "เกินกำหนด", "ชำระแล้ว", "ชำระล่าช้า"]
        status = st.selectbox("สถานะบิล", status_opts, key="ar2_status")
    if status == "ยังไม่ชำระ":
        filtered = filtered[filtered["Unpaid"]]
    elif status == "เกินกำหนด":
        filtered = filtered[filtered["OverdueFlag"]]
    elif status == "ชำระแล้ว":
        filtered = filtered[filtered["Paid"]]
    elif status == "ชำระล่าช้า":
        filtered = filtered[filtered["Status"] == "ชำระล่าช้า"]
    with f5:
        cls_present = [i for i in range(4) if (filtered["DebtorClass"] == i).any()]
        chosen_c = st.multiselect("ชั้นลูกหนี้ (TFRS 9)", [CLASS_NAMES[i] for i in cls_present],
                                  placeholder="ทั้งหมด", key="ar2_class")
    if chosen_c:
        idx = [CLASS_NAMES.index(c) for c in chosen_c]
        filtered = filtered[filtered["DebtorClass"].isin(idx)]
    with f6:
        search = st.text_input("ค้นหาลูกหนี้", placeholder="พิมพ์ชื่อลูกหนี้บางส่วน", key="ar2_search")
    if search.strip():
        filtered = filtered[filtered["Debtor"].str.contains(search.strip(), case=False, regex=False, na=False)]

    if filtered.empty:
        st.info("ไม่มีข้อมูลตามตัวกรองที่เลือก")
        return

    unpaid = filtered[filtered["Unpaid"]]
    paid = filtered[filtered["Paid"]]
    debtors = debtor_table(filtered, src)

    # ---------------- KPI ----------------
    with _safe("KPI"):
        outstanding = float(unpaid["Amount"].sum())
        overdue = float(filtered["OverdueAmt"].sum())
        npl_amt = float(unpaid.loc[unpaid["AgeIdx"] == 4, "Amount"].sum())
        npl_debtors = int(unpaid.loc[unpaid["AgeIdx"] == 4, "Debtor"].nunique())
        risky = debtors[(debtors["Class"].isin([1, 2])) & (debtors["Outstanding"] > 0)]
        on_time = float((~paid["IsLate"]).mean()) if not paid.empty else float("nan")
        late_paid = paid[paid["IsLate"]]
        avg_late = float(late_paid["DaysLate"].mean()) if not late_paid.empty else 0.0
        collect_days = _wavg((paid["PayDate"] - paid["InvDate"]).dt.days.astype("float64"), paid["Amount"])
        credit_avg = _wavg(paid["Credit"].astype("float64"), paid["Amount"])
        over_credit = _wavg(paid["DaysLate"].astype("float64"), paid["Amount"])

        def kpi(icon, bg, label, value, sub, color="#0F172A"):
            return (f'<div class="tc-kpi"><div class="tc-kpi-icon" style="background:{bg}">{icon}</div>'
                    f'<div><div class="tc-kpi-label">{esc(label)}</div>'
                    f'<div class="tc-kpi-value" style="color:{color}">{esc(value)}</div>'
                    f'<div class="tc-kpi-sub">{sub}</div></div></div>')

        cards = [
            kpi("💰", "#FFE8EC", "ยอดลูกหนี้คงค้าง", _money(outstanding),
                f"{len(unpaid):,} บิล · {unpaid['Debtor'].nunique():,} ราย"),
            kpi("⚠️", "#FDE9EC", "เกินกำหนดชำระ", _money(overdue),
                f"{(overdue / outstanding * 100 if outstanding else 0):.1f}% ของยอดคงค้าง · "
                f"{int(filtered['OverdueFlag'].sum()):,} บิล", "#C23B53" if overdue else "#0F172A"),
            kpi("⏳", "#FFF0E6", "ค้างเกิน 90 วัน (NPL)", _money(npl_amt),
                f"{npl_debtors:,} ราย", "#C23B53" if npl_amt else "#0F172A"),
            kpi("🚨", "#FFF3D6", "ลูกหนี้เสี่ยงที่ยังมียอดค้าง", f"{len(risky):,} ราย",
                f"เฝ้าติดตาม {int((risky['Class'] == 1).sum()):,} · ด้อยคุณภาพ {int((risky['Class'] == 2).sum()):,}"),
            kpi("✅", "#E7F6EC", "ชำระตรงเวลา",
                "—" if pd.isna(on_time) else f"{on_time * 100:.1f}%",
                f"จาก {len(paid):,} บิลที่ชำระแล้ว · บิลที่ช้าเฉลี่ย {avg_late:,.0f} วัน"),
            kpi("📅", "#EEF2FF", "เก็บเงินได้จริงเฉลี่ย",
                "—" if pd.isna(collect_days) else f"{collect_days:,.0f} วัน",
                ("หลังวางบิล (ถ่วงตามยอดเงิน)" if pd.isna(credit_avg)
                 else f"เครดิตเฉลี่ย {credit_avg:,.0f} วัน · เกินเครดิตเฉลี่ย {0 if pd.isna(over_credit) else over_credit:,.1f} วัน")),
        ]
        st.markdown('<div class="tc-kpi-grid">' + "".join(cards) + "</div>", unsafe_allow_html=True)

    # ---------------- สรุปสำหรับผู้บริหาร ----------------
    with _safe("สรุปสำหรับผู้บริหาร"):
        rows = []  # (โทน, ไอคอน, ข้อความ html)
        n_out_debtors = int(unpaid["Debtor"].nunique())
        if outstanding > 0:
            share = overdue / outstanding
            tone = "bad" if share >= 0.3 else ("warn" if overdue > 0 else "good")
            rows.append((tone, "💰",
                         f"ยอดลูกหนี้คงค้าง <b>{_money(outstanding)}</b> จาก {n_out_debtors:,} ราย · "
                         f"เกินกำหนดแล้ว <b>{_money(overdue)}</b> ({share * 100:.0f}% ของยอดคงค้าง)"))
        else:
            rows.append(("good", "✅", "ไม่มียอดลูกหนี้คงค้างตามตัวกรอง — ทุกบิลชำระแล้ว"))

        if npl_amt > 0:
            rows.append(("bad", "🚨",
                         f"ค้างเกิน 90 วัน (NPL) <b>{_money(npl_amt)}</b> จาก {npl_debtors:,} ราย — "
                         "ควรพิจารณาระงับการขายและตั้งค่าเผื่อหนี้สงสัยจะสูญ"))
        elif outstanding > 0:
            rows.append(("good", "🛡️", "ไม่มียอดค้างเกิน 90 วัน (ไม่มีหนี้ด้อยคุณภาพ)"))

        if outstanding > 0:
            per = unpaid.groupby("Debtor")["Amount"].sum().sort_values(ascending=False)
            top_n = min(5, len(per))
            top_share = float(per.head(top_n).sum() / outstanding)
            tone = "warn" if top_share >= 0.5 and len(per) > top_n else "info"
            rows.append((tone, "🎯",
                         f"ลูกหนี้ {top_n} รายแรกคิดเป็น <b>{top_share * 100:.0f}%</b> ของยอดคงค้าง · "
                         f"รายใหญ่สุด {esc(str(per.index[0])[:40])} <b>{_money(per.iloc[0])}</b>"
                         + (" — ความเสี่ยงกระจุกตัว ควรติดตามใกล้ชิด" if tone == "warn" else "")))

        if not paid.empty:
            tone = "good" if on_time >= 0.8 else ("warn" if on_time >= 0.5 else "bad")
            txt = f"ชำระตรงเวลา <b>{on_time * 100:.0f}%</b> ของบิลที่ชำระแล้ว ({len(paid):,} บิล)"
            if not pd.isna(collect_days):
                txt += f" · เก็บเงินได้จริงเฉลี่ย <b>{collect_days:,.0f} วัน</b> หลังวางบิล"
                if not pd.isna(credit_avg):
                    txt += f" (เครดิตเฉลี่ย {credit_avg:,.0f} วัน)"
            rows.append((tone, "⏱️", txt))

        # เทียบเดือนล่าสุดกับเดือนก่อน (ตามเดือนที่ชำระ ไม่นับเดือนปัจจุบันที่ยังไม่จบ)
        pm = paid[paid["PayDate"].notna()]
        pm = pm[pm["PayDate"].dt.to_period("M") < ref_date.to_period("M")] if pd.notna(ref_date) else pm
        if not pm.empty:
            by_m = pm.groupby(pm["PayDate"].dt.to_period("M")).agg(
                Late=("IsLate", "mean"), Bills=("Amount", "size"), Amt=("Amount", "sum")).sort_index()
            if len(by_m) >= 2 and by_m["Bills"].iloc[-1] >= 5 and by_m["Bills"].iloc[-2] >= 5:
                cur, prev = by_m.iloc[-1], by_m.iloc[-2]
                p_cur, p_prev = by_m.index[-1], by_m.index[-2]
                diff = (cur["Late"] - prev["Late"]) * 100
                if abs(diff) < 0.5:
                    trend, tone = "ทรงตัว", "info"
                elif diff > 0:
                    trend, tone = f'<span class="ar-up">▲ แย่ลง {abs(diff):.0f} จุด</span>', "warn"
                else:
                    trend, tone = f'<span class="ar-down">▼ ดีขึ้น {abs(diff):.0f} จุด</span>', "good"
                rows.append((tone, "📈",
                             f"เดือน {THAI_MONTHS[p_cur.month]} {str(p_cur.year + 543)[-2:]} ชำระช้า "
                             f"<b>{cur['Late'] * 100:.0f}%</b> ของ {int(cur['Bills']):,} บิลที่รับชำระ · "
                             f"เทียบ {THAI_MONTHS[p_prev.month]} {str(p_prev.year + 543)[-2:]} "
                             f"{prev['Late'] * 100:.0f}% → {trend}"))

        st.markdown(
            '<div class="ar-exec"><div class="ar-exec-head">📋 สรุปสำหรับผู้บริหาร</div>'
            + "".join(f'<div class="ar-exec-row {tone}"><span class="ic">{icon}</span><div>{txt}</div></div>'
                      for tone, icon, txt in rows)
            + "</div>",
            unsafe_allow_html=True,
        )

    # ---------------- ต้องลงมือ ----------------
    with _safe("ต้องลงมือ"):
        with _card("ar_actions"):
            st.markdown("#### ลูกหนี้ที่ต้องลงมือ (เรียงตามความเร่งด่วน)")
            st.caption("แต่ละรายอยู่ในกลุ่มตามบิลที่ค้างนานที่สุด (ยอดที่แสดง = ยอดค้างทั้งหมดของรายนั้น) · "
                       "จ่ายช้าเป็นประจำ = บิลที่ช้า ≥50% "
                       "และมีอย่างน้อย 3 บิล (จากชีตสรุปลูกหนี้ ถ้ามี)")
            worst_age = unpaid.groupby("Debtor")["AgeIdx"].max()
            out_by = unpaid.groupby("Debtor").agg(Out=("Amount", "sum"), Days=("DaysLate", "max"))
            out_by["Age"] = worst_age
            groups = [
                ("🔴 ค้างเกิน 90 วัน", "พิจารณาระงับการขาย / ตั้งสำรองหนี้สูญ / ส่งทนาย", "#C23B53",
                 out_by[out_by["Age"] == 4], [4]),
                ("🟠 ค้าง 31–90 วัน", "เร่งทวงถาม / จำกัดวงเงินเครดิต", "#E07B39",
                 out_by[out_by["Age"].isin([2, 3])], [2, 3]),
                ("🟡 ค้าง 1–30 วัน", "โทรเตือนก่อนกลายเป็นหนี้เสีย", "#C9961A",
                 out_by[out_by["Age"] == 1], [1]),
            ]
            cards_html = []
            for title, action, color, g, ages in groups:
                g = g.sort_values("Out", ascending=False)
                in_bucket = float(unpaid.loc[unpaid["AgeIdx"].isin(ages) & unpaid["Debtor"].isin(g.index), "Amount"].sum())
                items = "".join(
                    f'<div class="ar-act-item"><span class="n" title="{esc(str(n))}">{esc(str(n))}</span>'
                    f'<span class="v">{_money(r["Out"])}<small>ค้าง {int(r["Days"]):,} วัน</small></span></div>'
                    for n, r in g.head(5).iterrows()
                )
                more = f'<div class="ar-act-more">และอีก {len(g) - 5:,} ราย</div>' if len(g) > 5 else ""
                body = (f'<div class="ar-act-num">{len(g):,} <span>ราย · ยอดค้างรวม {_money(g["Out"].sum())}</span></div>'
                        f'<div class="ar-act-do">ในนั้นเป็นบิลช่วงนี้ {_money(in_bucket)}</div>'
                        f'<div class="ar-act-list">{items}</div>{more}') if not g.empty \
                    else '<div class="ar-act-empty">✅ ไม่มีลูกหนี้ในกลุ่มนี้</div>'
                cards_html.append(f'<div class="ar-act" style="--c:{color}"><div class="ar-act-title">{title}</div>'
                                  f'<div class="ar-act-do">👉 {action}</div>{body}</div>')

            habit = debtors[(debtors["LateRate"] >= 0.5) & (debtors["Bills"] >= 3)].sort_values(
                ["LateRate", "Bills"], ascending=False)
            items = "".join(
                f'<div class="ar-act-item"><span class="n" title="{esc(str(r["Debtor"]))}">{esc(str(r["Debtor"]))}</span>'
                f'<span class="v">{r["LateRate"] * 100:.0f}% ช้า<small>{int(r["LateBills"]):,}/{int(r["Bills"]):,} บิล</small></span></div>'
                for r in habit.head(5).to_dict("records")
            )
            more = f'<div class="ar-act-more">และอีก {len(habit) - 5:,} ราย</div>' if len(habit) > 5 else ""
            body = (f'<div class="ar-act-num">{len(habit):,} <span>ราย</span></div>'
                    f'<div class="ar-act-list">{items}</div>{more}') if not habit.empty \
                else '<div class="ar-act-empty">✅ ไม่มีลูกหนี้ที่จ่ายช้าเป็นประจำ</div>'
            cards_html.append('<div class="ar-act" style="--c:#7B5BB5"><div class="ar-act-title">🔁 จ่ายช้าเป็นประจำ</div>'
                              f'<div class="ar-act-do">👉 ทบทวนเครดิตเทอม / ขอเงื่อนไขชำระก่อนส่งของ</div>{body}</div>')
            st.markdown('<div class="ar-act-grid">' + "".join(cards_html) + "</div>", unsafe_allow_html=True)

    # ---------------- Aging + ชั้นลูกหนี้ ----------------
    with _safe("Aging + ชั้นลูกหนี้"):
        left, right = st.columns([1.1, 0.9], gap="medium")
        with left:
            with _card("ar_aging"):
                st.markdown("#### อายุของยอดลูกหนี้คงค้าง (Aging)")
                st.caption("เฉพาะบิลที่ยังไม่ชำระ · จัดกลุ่มจากตัวเลขวันช้าตามเกณฑ์ในชีตเงื่อนไข")
                ag = (unpaid.groupby("AgeIdx")
                      .agg(Amt=("Amount", "sum"), Bills=("Amount", "size"), Debtors=("Debtor", "nunique"))
                      .reindex(range(5), fill_value=0))
                if ag["Amt"].sum() <= 0:
                    st.info("ไม่มียอดคงค้างตามตัวกรอง")
                else:
                    total = float(ag["Amt"].sum())
                    fig = go.Figure(go.Bar(
                        x=AGING_LABELS, y=ag["Amt"], marker_color=AGING_COLORS,
                        text=[f"{_money(v)}<br><span style='font-size:11px;color:#94A3B8'>"
                              f"{(v / total * 100):.0f}%</span>" for v in ag["Amt"]],
                        textposition="outside", cliponaxis=False,
                        customdata=[[int(b), int(d)] for b, d in zip(ag["Bills"], ag["Debtors"])],
                        hovertemplate=("%{x}<br>ยอดคงค้าง: ฿%{y:,.0f}<br>%{customdata[0]:,} บิล · "
                                       "%{customdata[1]:,} ราย<extra></extra>"),
                    ))
                    fig.update_layout(
                        height=350, margin=dict(l=10, r=10, t=30, b=20), showlegend=False,
                        yaxis=dict(title="บาท", tickformat=",.0f", range=[0, float(ag["Amt"].max()) * 1.28]),
                        xaxis=dict(title=""), bargap=0.35,
                    )
                    _style(fig)
                    st.plotly_chart(fig, width="stretch", config=PLOT_CONFIG)

        with right:
            with _card("ar_class"):
                st.markdown("#### ยอดคงค้างตามชั้นลูกหนี้ (TFRS 9)")
                st.caption("ตามชั้นของแต่ละบิลในไฟล์ · ใช้ประเมินการตั้งค่าเผื่อหนี้สงสัยจะสูญ")
                cl = (unpaid.groupby("ClassIdx")
                      .agg(Amt=("Amount", "sum"), Debtors=("Debtor", "nunique"))
                      .reindex(range(4), fill_value=0))
                cl = cl[cl["Amt"] > 0]
                if cl.empty:
                    st.info("ไม่มียอดคงค้างตามตัวกรอง")
                else:
                    total = float(cl["Amt"].sum())
                    fig = go.Figure(go.Pie(
                        labels=[CLASS_NAMES[i] for i in cl.index], values=cl["Amt"], hole=0.66, sort=False,
                        marker=dict(colors=[CLASS_COLORS[i] for i in cl.index], line=dict(color="white", width=3)),
                        textinfo="none",
                        hovertemplate="%{label}<br>฿%{value:,.0f}<br>%{percent}<extra></extra>",
                    ))
                    fig.update_layout(
                        height=210, showlegend=False, margin=dict(l=0, r=0, t=4, b=4),
                        annotations=[dict(text=f"<span style='font-size:11px;color:#64748B'>ยอดคงค้าง</span>"
                                               f"<br><b style='font-size:16px;color:#1E293B'>{_money(total)}</b>",
                                          x=0.5, y=0.5, showarrow=False)],
                    )
                    _style(fig)
                    st.plotly_chart(fig, width="stretch", config=PLOT_CONFIG)
                    rules = {_class_idx(n): r for n, r, _m in class_crit}
                    st.markdown(
                        "".join(
                            f'<div class="ar-leg-row"><span class="sw" style="background:{CLASS_COLORS[i]}"></span>'
                            f'<div><div class="nm">{esc(CLASS_NAMES[i])}</div>'
                            f'<div class="rule">{esc(rules.get(i, CLASS_DEFAULT_RULE[i]))}</div></div>'
                            f'<div class="val">{_money(r["Amt"])}<small>{r["Amt"] / total * 100:.1f}% · '
                            f'{int(r["Debtors"]):,} ราย</small></div></div>'
                            for i, r in cl.iterrows()
                        ),
                        unsafe_allow_html=True,
                    )

    # ---------------- การกระจุกตัว ----------------
    with _safe("การกระจุกตัว"):
        with _card("ar_pareto"):
            st.markdown("#### การกระจุกตัวของยอดคงค้าง (10 รายใหญ่)")
            st.caption("แท่ง = ยอดคงค้างของลูกหนี้แต่ละราย · เส้น = % สะสมของยอดคงค้างทั้งหมด")
            per = unpaid.groupby("Debtor")["Amount"].sum().sort_values(ascending=False)
            if per.sum() <= 0:
                st.info("ไม่มียอดคงค้างตามตัวกรอง")
            else:
                top = per.head(10)
                cum = top.cumsum() / per.sum() * 100
                names = [str(n) if len(str(n)) <= 22 else str(n)[:20] + "…" for n in top.index]
                cls = debtors.set_index("Debtor")["Class"].reindex(top.index).fillna(3).astype(int)
                fig = go.Figure()
                fig.add_trace(go.Bar(
                    x=names, y=top.values, marker_color=[CLASS_COLORS[i] for i in cls], name="ยอดคงค้าง",
                    customdata=[[str(n), CLASS_NAMES[i]] for n, i in zip(top.index, cls)],
                    hovertemplate="%{customdata[0]}<br>฿%{y:,.0f}<br>%{customdata[1]}<extra></extra>",
                ))
                fig.add_trace(go.Scatter(
                    x=names, y=cum.values, yaxis="y2", mode="lines+markers+text", name="% สะสม",
                    line=dict(color="#7B5BB5", width=2.5), marker=dict(size=7),
                    text=[f"{v:.0f}%" for v in cum.values], textposition="top center",
                    textfont=dict(color="#7B5BB5", size=11),
                    hovertemplate="สะสม %{y:.1f}%<extra></extra>",
                ))
                fig.update_layout(
                    height=360, margin=dict(l=10, r=10, t=20, b=20), showlegend=False, bargap=0.3,
                    xaxis=dict(title="", type="category", tickangle=-30),
                    yaxis=dict(title="บาท", tickformat=",.0f"),
                    yaxis2=dict(overlaying="y", side="right", range=[0, 110], ticksuffix="%", showgrid=False),
                )
                _style(fig)
                st.plotly_chart(fig, width="stretch", config=PLOT_CONFIG)
                st.caption(f"สีแท่ง = ชั้นลูกหนี้ · ทั้งหมด {len(per):,} รายที่มียอดค้าง")


    # ---------------- แนวโน้มรายเดือน ----------------
    with _safe("แนวโน้มรายเดือน"):
        with _card("ar_trend"):
            st.markdown("#### ยอดวางบิลรายเดือน และสัดส่วนบิลที่ชำระล่าช้า")
            st.caption("แท่ง = ยอดวางบิลแยกตามสถานะปัจจุบัน · เส้น = % บิลที่ชำระช้ากว่ากำหนด (จากบิลที่ชำระแล้วของเดือนนั้น)")
            t = filtered[filtered["InvDate"].notna()]
            if t.empty:
                st.info("ไม่มีข้อมูลวันที่วางบิล")
            else:
                t = t.assign(_P=t["InvDate"].dt.to_period("M"))
                g = t.groupby("_P").apply(lambda x: pd.Series({
                    "ชำระตรงเวลา": x.loc[x["Status"] == "ชำระตรงเวลา", "Amount"].sum(),
                    "ชำระล่าช้า": x.loc[x["Status"] == "ชำระล่าช้า", "Amount"].sum(),
                    "ยังไม่ถึงกำหนด": x.loc[x["Status"] == "ยังไม่ถึงกำหนด", "Amount"].sum(),
                    "เกินกำหนด": x.loc[x["Status"] == "เกินกำหนด", "Amount"].sum(),
                    "LateRate": (x.loc[x["Paid"], "IsLate"].mean() * 100) if x["Paid"].any() else float("nan"),
                })).sort_index()
                labels = [f"{THAI_MONTHS[p.month]} {str(p.year + 543)[-2:]}" for p in g.index]
                fig = go.Figure()
                for s in ["ชำระตรงเวลา", "ชำระล่าช้า", "ยังไม่ถึงกำหนด", "เกินกำหนด"]:
                    fig.add_trace(go.Bar(name=s, x=labels, y=g[s], marker_color=STATUS_COLORS[s],
                                         hovertemplate=f"{s}: ฿%{{y:,.0f}}<extra></extra>"))
                fig.add_trace(go.Scatter(
                    name="% บิลชำระล่าช้า", x=labels, y=g["LateRate"], yaxis="y2", mode="lines+markers+text",
                    line=dict(color="#7B5BB5", width=2.5, shape="spline", smoothing=0.5),
                    marker=dict(size=7, line=dict(color="white", width=1.5)),
                    text=[f"{v:.0f}%" if pd.notna(v) else "" for v in g["LateRate"]], textposition="top center",
                    textfont=dict(color="#7B5BB5", size=11),
                    hovertemplate="% บิลชำระล่าช้า: %{y:.1f}%<extra></extra>",
                ))
                fig.update_layout(
                    barmode="stack", height=400, margin=dict(l=10, r=10, t=20, b=20),
                    legend=dict(orientation="h", y=1.12, x=0), hovermode="x unified", bargap=0.3,
                    xaxis=dict(title="", type="category"),
                    yaxis=dict(title="ยอดวางบิล (บาท)", tickformat=",.0f"),
                    yaxis2=dict(title="% ชำระล่าช้า", overlaying="y", side="right", range=[0, 110],
                                showgrid=False, ticksuffix="%"),
                )
                _style(fig)
                st.plotly_chart(fig, width="stretch", config=PLOT_CONFIG)

    # ---------------- สาขา + เครดิต ----------------
    with _safe("สาขา + เครดิต"):
        _branch_pre = (unpaid.groupby("Branch")
                       .agg(Out=("Amount", "sum"), Over=("OverdueAmt", "sum")).reset_index())
        _branch_pre = _branch_pre[_branch_pre["Out"] > 0]
        _credit_pre = paid[paid["Credit"].notna()]
        n_branch_bars = min(len(_branch_pre), 12)
        n_credit_bars = _credit_pre["Credit"].nunique() if not _credit_pre.empty else 0
        ar_row_height = max(320, 36 * max(n_branch_bars, n_credit_bars) + 110)

        left, right = st.columns([1, 1], gap="medium")
        with left:
            with _card("ar_branch"):
                st.markdown("#### ยอดคงค้างตามสาขา")
                st.caption("แยกส่วนที่ยังไม่ถึงกำหนด และส่วนที่เกินกำหนดแล้ว · เรียงตามยอดเกินกำหนด")
                b = (unpaid.groupby("Branch")
                     .agg(Out=("Amount", "sum"), Over=("OverdueAmt", "sum")).reset_index())
                b = b[b["Out"] > 0]
                if b.empty:
                    st.info("ไม่มียอดคงค้างตามตัวกรอง")
                else:
                    b["NotDue"] = b["Out"] - b["Over"]
                    b = b.sort_values(["Over", "Out"], ascending=False).head(12).iloc[::-1]
                    fig = go.Figure()
                    fig.add_trace(go.Bar(name="ยังไม่ถึงกำหนด", y=b["Branch"], x=b["NotDue"], orientation="h",
                                         marker_color=STATUS_COLORS["ยังไม่ถึงกำหนด"],
                                         hovertemplate="%{y}<br>ยังไม่ถึงกำหนด: ฿%{x:,.0f}<extra></extra>"))
                    fig.add_trace(go.Bar(name="เกินกำหนด", y=b["Branch"], x=b["Over"], orientation="h",
                                         marker_color=STATUS_COLORS["เกินกำหนด"],
                                         hovertemplate="%{y}<br>เกินกำหนด: ฿%{x:,.0f}<extra></extra>"))
                    fig.add_trace(go.Scatter(
                        y=b["Branch"], x=b["Out"], mode="text", showlegend=False, hoverinfo="skip",
                        text=[f"  {_money(o)} · เกิน {(v / o * 100 if o else 0):.0f}%" for o, v in zip(b["Out"], b["Over"])],
                        textposition="middle right", textfont=dict(color="#334155", size=12),
                    ))
                    fig.update_layout(
                        barmode="stack", height=ar_row_height, margin=dict(l=10, r=20, t=20, b=30),
                        legend=dict(orientation="h", y=1.1, x=0), bargap=0.35,
                        xaxis=dict(title="บาท", tickformat=",.0f", range=[0, float(b["Out"].max()) * 1.45]),
                        yaxis=dict(title="", automargin=True),
                    )
                    _style(fig)
                    st.plotly_chart(fig, width="stretch", config=PLOT_CONFIG)

        with right:
            with _card("ar_credit"):
                st.markdown("#### พฤติกรรมการชำระตามระยะเวลาเครดิต")
                st.caption("% บิลที่ชำระช้ากว่ากำหนด แยกตามเครดิตที่ให้ลูกค้า · ใช้ทบทวนนโยบายเครดิต")
                cr = paid[paid["Credit"].notna()]
                if cr.empty:
                    st.info("ไม่มีข้อมูลระยะเวลาเครดิต")
                else:
                    c = (cr.groupby("Credit")
                         .agg(Bills=("Amount", "size"), Late=("IsLate", "mean"),
                              AvgLate=("DaysLate", lambda s: float(s[s > 0].mean()) if (s > 0).any() else 0.0))
                         .reset_index().sort_values("Credit"))
                    labels = [f"{int(v)} วัน" for v in c["Credit"]]
                    colors = ["#E0566C" if v >= 0.5 else ("#F6AE6B" if v >= 0.25 else "#86CFA3") for v in c["Late"]]
                    fig = go.Figure(go.Bar(
                        x=labels, y=c["Late"] * 100, marker_color=colors,
                        text=[f"{v * 100:.0f}%" for v in c["Late"]], textposition="outside", cliponaxis=False,
                        customdata=[[int(n), float(a)] for n, a in zip(c["Bills"], c["AvgLate"])],
                        hovertemplate=("เครดิต %{x}<br>ชำระช้า: %{y:.1f}% ของ %{customdata[0]:,} บิล"
                                       "<br>บิลที่ช้า ช้าเฉลี่ย %{customdata[1]:,.0f} วัน<extra></extra>"),
                    ))
                    fig.update_layout(
                        height=ar_row_height, margin=dict(l=10, r=10, t=20, b=20), showlegend=False,
                        xaxis=dict(title="ระยะเวลาเครดิต", type="category"),
                        yaxis=dict(title="% บิลที่ชำระช้า", range=[0, 112], ticksuffix="%"), bargap=0.15,
                    )
                    _style(fig)
                    st.plotly_chart(fig, width="stretch", config=PLOT_CONFIG)

    # ---------------- ลูกหนี้ที่ควรติดตาม ----------------
    with _safe("ลูกหนี้ที่ควรติดตาม"):
        with _card("ar_top"):
            st.markdown("#### ลูกหนี้ที่ควรติดตามก่อน")
            o1, o2, o3 = st.columns([1.4, 0.7, 1.2])
            sort_opts = {"ยอดเกินกำหนดมากสุด": ("OverdueAmt", False), "ยอดคงค้างมากสุด": ("Outstanding", False),
                         "ค้างนานที่สุด": ("MaxUnpaidDelay", False), "% บิลที่ช้ามากสุด": ("LateRate", False),
                         "วันช้าสูงสุด (ทั้งช่วง)": ("MaxDelay", False)}
            with o1:
                sort_label = st.selectbox("เรียงตาม", list(sort_opts), key="ar2_top_sort")
            with o2:
                top_n = st.selectbox("จำนวน", [10, 20, 30, 50, 100], index=1, key="ar2_top_n")
            with o3:
                st.markdown("<div style='height:1.75rem'></div>", unsafe_allow_html=True)
                only_open = st.checkbox("เฉพาะลูกหนี้ที่ยังมียอดค้าง", value=True, key="ar2_top_open")
            tb = debtors[debtors["Outstanding"] > 0] if only_open else debtors
            col, asc = sort_opts[sort_label]
            tb = tb.sort_values([col, "Outstanding"], ascending=[asc, False], na_position="last").reset_index(drop=True)
            st.caption(
                f"แสดง {min(top_n, len(tb)):,} จาก {len(tb):,} ราย · ยอดคงค้าง/เกินกำหนด/ค้างนานสุด เปลี่ยนตามตัวกรอง · "
                "จำนวนบิล / บิลที่ช้า / % / วันช้าสูงสุด / ชั้นลูกหนี้ มาจากชีตสรุปลูกหนี้"
                + ("" if src is not None else " (ไม่พบชีตสรุป — คำนวณจากรายบิลแทน)")
            )
            max_out = float(tb["Outstanding"].max() or 1) if not tb.empty else 1.0
            head = ["#", "ลูกหนี้", "สาขา", "ยอดคงค้าง", "เกินกำหนด", "ค้างนานสุด", "บิล (ช้า/ทั้งหมด)",
                    "% บิลที่ช้า", "วันช้าสูงสุด", "กลุ่มอายุหนี้", "ชั้นลูกหนี้"]
            body = []
            for i, r in enumerate(tb.head(top_n).to_dict("records"), 1):
                age_i = _aging_idx_from_text(r["AgingText"])
                age_i = r["AgingIdx"] if age_i is None else age_i
                late_rate = r["LateRate"]
                late_txt = "—" if pd.isna(late_rate) else f"{late_rate * 100:.0f}%"
                wait = int(r["MaxUnpaidDelay"])
                wait_txt = f"{wait:,} วัน" if wait > 0 else "—"
                body.append(
                    "<tr>"
                    f'<td class="tc-rank">{i}</td>'
                    f"<td><b>{_name_html(r['Debtor'])}</b></td>"
                    f'<td class="tc-muted">{esc(r["Branch"])}</td>'
                    f'<td class="tc-num"><b>{_money_full(r["Outstanding"])}</b>'
                    f'<div class="tc-bar"><span style="width:{r["Outstanding"] / max_out * 100:.1f}%;'
                    f'background:#F07C8C"></span></div></td>'
                    f'<td class="tc-num {"ar-over" if r["OverdueAmt"] > 0 else "tc-muted"}">'
                    f'{_money_full(r["OverdueAmt"]) if r["OverdueAmt"] > 0 else "—"}</td>'
                    f'<td class="tc-num {"ar-over" if r["MaxUnpaidDelay"] > 0 else "tc-muted"}">'
                    f'{wait_txt}</td>'
                    f'<td class="tc-num">{int(r["LateBills"]):,} / {int(r["Bills"]):,}</td>'
                    f'<td class="tc-num">{late_txt}</td>'
                    f'<td class="tc-num">{int(r["MaxDelay"]):,} วัน</td>'
                    f"<td>{_badge(r['AgingText'] or AGING_KEYS[int(age_i)], AGING_COLORS[int(age_i)])}</td>"
                    f"<td>{_class_badge(r['Class'])}</td>"
                    "</tr>"
                )
            st.markdown(_table_html(head, body, right_cols={3, 4, 5, 6, 7, 8}, max_height=520), unsafe_allow_html=True)
            export = pd.DataFrame({
                "ลูกหนี้": tb["Debtor"], "สาขา": tb["Branch"], "ยอดคงค้าง": tb["Outstanding"],
                "เกินกำหนด": tb["OverdueAmt"], "ค้างนานสุด (วัน)": tb["MaxUnpaidDelay"],
                "จำนวนบิลทั้งหมด": tb["Bills"], "จำนวนบิลที่ช้า": tb["LateBills"],
                "% บิลที่ช้า": (tb["LateRate"] * 100).round(1), "วันช้าสูงสุด": tb["MaxDelay"],
                "กลุ่มอายุหนี้": tb["AgingText"], "ชั้นลูกหนี้": tb["Class"].map(lambda i: CLASS_NAMES[int(i)]),
            })
            st.download_button("⬇️ ดาวน์โหลดรายชื่อลูกหนี้ (CSV)", export.to_csv(index=False).encode("utf-8-sig"),
                               file_name="ar_debtors.csv", mime="text/csv", key="ar2_top_dl")

    # ---------------- รายละเอียดรายบิล ----------------
    with _safe("รายละเอียดรายบิล"):
        with _card("ar_bills"):
            st.markdown("#### รายละเอียดรายบิล")
            debtor_opts = ["ทุกลูกหนี้"] + debtors.sort_values("Outstanding", ascending=False)["Debtor"].tolist()
            b1, b2 = st.columns([2.2, 1])
            with b1:
                pick = st.selectbox("ลูกหนี้", debtor_opts, key=_wkey("ar2_bill_debtor", debtor_opts))
            with b2:
                bill_status = st.selectbox("สถานะ", ["ทั้งหมด", "ยังไม่ถึงกำหนด", "เกินกำหนด", "ชำระตรงเวลา", "ชำระล่าช้า"],
                                           key="ar2_bill_status")
            bills = filtered if pick == "ทุกลูกหนี้" else filtered[filtered["Debtor"] == pick]
            if bill_status != "ทั้งหมด":
                bills = bills[bills["Status"] == bill_status]
            bills = bills.sort_values(["Unpaid", "DaysLate", "Amount"], ascending=[False, False, False])
            st.caption(f"{len(bills):,} บิล · ยอดรวม {_money_full(bills['Amount'].sum())} · "
                       "เรียงบิลค้างชำระก่อน แล้วตามวันช้า · แสดงสูงสุด 500 บิล")
            head = ["เลขที่ใบวางบิล", "ลูกหนี้", "สาขา", "เครดิต", "วันที่วางบิล", "ครบกำหนด", "วันที่จ่าย",
                    "จำนวนเงิน", "วันช้า", "สถานะ", "ชั้นลูกหนี้"]
            body = []
            for r in bills.head(500).to_dict("records"):
                credit_txt = "—" if pd.isna(r["Credit"]) else f"{int(r['Credit'])} วัน"
                body.append(
                    "<tr>"
                    f'<td class="tc-mono">{esc(r["Bill"])}</td>'
                    f"<td>{_name_html(r['Debtor'])}</td>"
                    f'<td class="tc-muted">{esc(r["Branch"])}</td>'
                    f'<td class="tc-num tc-muted">{credit_txt}</td>'
                    f'<td class="tc-muted">{_thai_date(r["InvDate"])}</td>'
                    f'<td class="tc-muted">{_thai_date(r["DueDate"])}</td>'
                    f'<td class="tc-muted">{_thai_date(r["PayDate"])}</td>'
                    f'<td class="tc-num"><b>{_money_full(r["Amount"])}</b></td>'
                    f'<td class="tc-num {"ar-over" if r["DaysLate"] > 0 else "tc-muted"}">{int(r["DaysLate"]):,}</td>'
                    f"<td>{_badge(r['Status'], STATUS_COLORS[r['Status']])}</td>"
                    f"<td>{_class_badge(r['ClassIdx'])}</td>"
                    "</tr>"
                )
            st.markdown(_table_html(head, body, right_cols={3, 7, 8}, max_height=460), unsafe_allow_html=True)
            export = pd.DataFrame({
                "เลขที่ใบวางบิล": bills["Bill"], "ลูกหนี้": bills["Debtor"], "สาขา": bills["Branch"],
                "เครดิต (วัน)": bills["Credit"], "วันที่วางบิล": bills["InvDate"].dt.strftime("%d/%m/%Y"),
                "ครบกำหนด": bills["DueDate"].dt.strftime("%d/%m/%Y"), "วันที่จ่าย": bills["PayDate"].dt.strftime("%d/%m/%Y"),
                "จำนวนเงิน": bills["Amount"], "วันช้า": bills["DaysLate"], "สถานะ": bills["Status"],
                "ชั้นลูกหนี้": bills["ClassIdx"].map(lambda i: CLASS_NAMES[int(i)]),
            })
            st.download_button(f"⬇️ ดาวน์โหลด {len(bills):,} บิล (CSV)", export.to_csv(index=False).encode("utf-8-sig"),
                               file_name="ar_bills.csv", mime="text/csv", key="ar2_bill_dl")

    # ---------------- เกณฑ์อ้างอิง ----------------
    with _safe("เกณฑ์อ้างอิง"):
        with st.expander("📘 เกณฑ์ช่วงอายุหนี้ และการจัดชั้นลูกหนี้ (TFRS 9) — จากชีตเงื่อนไข"):
            def crit_table(title, rows):
                if not rows:
                    return ""
                return (f"<p><b>{esc(title)}</b></p><table class='ar-crit'><tr><th>กลุ่ม</th><th>เกณฑ์</th>"
                        "<th>ความหมาย</th></tr>"
                        + "".join(f"<tr><td>{esc(a)}</td><td>{esc(b)}</td><td>{esc(c)}</td></tr>" for a, b, c in rows)
                        + "</table>")
            html = crit_table("ช่วงอายุลูกหนี้ (Aging)", aging_crit) + crit_table("ชั้นลูกหนี้ (TFRS 9)", class_crit)
            st.markdown(html or "ไม่พบชีตเงื่อนไขในไฟล์", unsafe_allow_html=True)

        st.markdown(
            f'<div class="tc-source">ข้อมูลจาก {esc(path.name)} · ชีต {esc(str(detail_sheet))}'
            + (f" + {esc(str(summary_sheet))}" if summary_sheet else "")
            + f" · {len(ar):,} บิลทั้งไฟล์ · แสดง {len(filtered):,} บิลตามตัวกรอง</div>",
            unsafe_allow_html=True,
        )